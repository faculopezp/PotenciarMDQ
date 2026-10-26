"""Muestras de fuentes nuevas -> reports/fuentes/ (el CI las commitea).

El entorno de desarrollo no llega a *.gob.ar: para diseñar un parser hace falta ver el archivo
real. Este paso deja en el repo el encabezado, algunas filas y las filas del foco de cada CSV, y el
texto de cada PDF (solo las páginas que nombran al foco, más la primera).
"""

from __future__ import annotations

import csv
import io
import logging
from pathlib import Path

from potenciar_mdq.settings import RAW, REPORTS

log = logging.getLogger(__name__)
FOCO = ("mar del plata", "mardel", "general pueyrredon", "sazm")


def _csv(path: Path, n: int = 8, n_foco: int = 40) -> list[str]:
    raw = path.read_bytes()
    texto = next((raw.decode(enc) for enc in ("utf-8-sig", "latin-1") if _decodifica(raw, enc)), "")
    filas = list(csv.reader(io.StringIO(texto)))
    foco = [f for f in filas[1:] if any(k in ",".join(f).lower() for k in FOCO)]
    out = [f"### `{path.relative_to(RAW)}` — {len(filas) - 1} filas, {len(foco)} del foco", "```"]
    out += [",".join(f) for f in filas[: n + 1]]
    out += ["... (filas del foco)"] + [",".join(f) for f in foco[:n_foco]] + ["```", ""]
    return out


def _decodifica(raw: bytes, enc: str) -> bool:
    try:
        raw.decode(enc)
        return True
    except UnicodeDecodeError:
        return False


def _pdf(path: Path, dest: Path) -> str:
    import pdfplumber

    partes = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            texto = page.extract_text() or ""
            if i == 0 or any(k in texto.lower() for k in FOCO):
                partes.append(f"=== página {i + 1} de {len(pdf.pages)} ===\n{texto}")
    dest.write_text("\n\n".join(partes), encoding="utf-8")
    return f"- `{path.name}` → `{dest.relative_to(REPORTS)}` ({len(partes)} páginas)"


def _texto_csv(texto: str, nombre: str, n: int = 8, n_foco: int = 40) -> list[str]:
    lineas = texto.splitlines()
    foco = [x for x in lineas[1:] if any(k in x.lower() for k in FOCO)]
    out = [f"### `{nombre}` — {len(lineas) - 1} filas, {len(foco)} del foco", "```"]
    return out + lineas[: n + 1] + ["... (filas del foco)"] + foco[:n_foco] + ["```", ""]


def _zip(path: Path) -> list[str]:
    import zipfile

    out = [f"### `{path.relative_to(RAW)}` (zip)"]
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            out.append(f"- `{info.filename}` — {info.file_size:,} bytes")
        for info in z.infolist():
            if info.filename.lower().endswith((".csv", ".txt")):
                raw = z.read(info)
                texto = next((raw.decode(e) for e in ("utf-8-sig", "latin-1")
                              if _decodifica(raw, e)), "")
                out += _texto_csv(texto, f"{path.name}/{info.filename}")
    return out + [""]


def _xlsx(path: Path, n: int = 12, n_foco: int = 60) -> list[str]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = [f"### `{path.relative_to(RAW)}` (xlsx) — hojas: {', '.join(wb.sheetnames)}"]
    for ws in wb.worksheets:
        filas = [["" if v is None else str(v) for v in fila]
                 for fila in ws.iter_rows(values_only=True)]
        filas = [f for f in filas if any(f)]
        foco = [f for f in filas if any(k in " ".join(f).lower() for k in FOCO)]
        out += [f"#### hoja `{ws.title}` — {len(filas)} filas, {len(foco)} del foco", "```"]
        out += [" | ".join(f).rstrip(" |") for f in filas[:n]]
        out += ["... (filas del foco)"] + [" | ".join(f).rstrip(" |") for f in foco[:n_foco]]
        out += ["```", ""]
    return out


def _listado(sid: str) -> list[str]:
    archivos = sorted(p for p in (RAW / sid).rglob("*") if p.is_file())
    return [f"- `{p.relative_to(RAW)}` — {p.stat().st_size:,} bytes" for p in archivos] + [""]


def run(out: Path = REPORTS / "fuentes") -> Path:
    out.mkdir(parents=True, exist_ok=True)
    lineas = ["# Muestras de fuentes nuevas", ""]
    (out / "cnrt").mkdir(exist_ok=True)
    for sid in ("yvera_terrestre", "cnrt_ferroviario"):
        lineas += [f"## {sid}", "", "Archivos descargados:"] + _listado(sid)
        for p in sorted((RAW / sid).rglob("*")):
            suf = p.suffix.lower()
            try:
                if suf == ".csv":
                    lineas += _csv(p)
                elif suf == ".zip":
                    lineas += _zip(p)
                elif suf in (".xlsx", ".xlsm"):
                    lineas += _xlsx(p)
                elif suf == ".pdf":
                    lineas.append(_pdf(p, out / "cnrt" / f"{p.stem}.txt"))
            except Exception as exc:  # noqa: BLE001 - muestra, no corta
                lineas.append(f"- ❌ {p.name}: {type(exc).__name__}: {exc}")
    dest = out / "muestras.md"
    dest.write_text("\n".join(lineas), encoding="utf-8")
    return dest
